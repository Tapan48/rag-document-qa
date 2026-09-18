from dataclasses import dataclass
from pathlib import Path

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class ExtractionError(Exception):
    """Raised when a document's text cannot be extracted."""


@dataclass
class TextUnit:
    text: str
    source: dict


def extract_units(path: Path, extension: str) -> list[TextUnit]:
    extension = extension.lower()
    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".docx":
        return _extract_docx(path)
    if extension == ".txt":
        return _extract_txt(path)
    raise ExtractionError(f"Unsupported file extension: {extension}")


def _extract_pdf(path: Path) -> list[TextUnit]:
    try:
        reader = PdfReader(path)
    except PdfReadError as exc:
        raise ExtractionError("Could not read PDF (corrupt or invalid)") from exc

    if reader.is_encrypted:
        raise ExtractionError("Encrypted PDFs are not supported")

    units = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            units.append(TextUnit(text=text, source={"kind": "page", "page": page_number}))

    if not units:
        raise ExtractionError(
            "No extractable text found (scanned PDFs require OCR, which is unsupported)"
        )
    return units


def _extract_docx(path: Path) -> list[TextUnit]:
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError("Could not read DOCX (corrupt or invalid)") from exc

    units = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            units.append(TextUnit(text=text, source={"kind": "paragraph", "paragraph": index}))

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            row_text = " | ".join(cell.text.strip() for cell in row.cells).strip()
            if row_text:
                units.append(
                    TextUnit(
                        text=row_text,
                        source={"kind": "table_row", "table": table_index, "row": row_index},
                    )
                )

    if not units:
        raise ExtractionError("No extractable text found in DOCX")
    return units


def _extract_txt(path: Path) -> list[TextUnit]:
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ExtractionError("File is not valid UTF-8 text") from exc

    units = [
        TextUnit(text=line.strip(), source={"kind": "line", "line": line_number})
        for line_number, line in enumerate(content.splitlines(), start=1)
        if line.strip()
    ]
    if not units:
        raise ExtractionError("No extractable text found in TXT file")
    return units
