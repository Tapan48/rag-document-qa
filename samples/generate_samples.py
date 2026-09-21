"""Generates the three fictional sample documents used for demos and evaluation.

Single source of truth: the content is defined once below, then rendered into
both the human-editable Markdown source (samples/source/) and the actual
PDF/DOCX/TXT files uploaded to the app (samples/). Re-run this script after
editing the content constants to regenerate all six files.

    python samples/generate_samples.py

Everything here is fictional -- "Solstice Home Hub" and "Solstice Devices
Inc." are not real products or companies.
"""

from pathlib import Path

import docx
from fpdf import FPDF

SAMPLES_DIR = Path(__file__).parent
SOURCE_DIR = SAMPLES_DIR / "source"

# ---------------------------------------------------------------------------
# Product guide (PDF) -- page 1: overview/specs/box contents, page 2: setup/LEDs
# ---------------------------------------------------------------------------

PRODUCT_GUIDE_PAGE_1 = """Solstice Home Hub SH-100 -- Product Guide

Overview
The Solstice Home Hub SH-100, made by Solstice Devices Inc., is a smart home
hub that connects and controls compatible smart devices from a single app.

Wireless connectivity
The SH-100 supports Wi-Fi 6, Bluetooth 5.2, Zigbee 3.0, and Thread. It does
not support Apple HomeKit.

Device capacity
The hub supports up to 150 connected smart devices at once. Connecting more
than 150 devices is not supported and can cause connectivity issues; see the
Troubleshooting Guide for what to do if you have too many devices connected.

Power
The SH-100 is powered by USB-C using the included 9W adapter. It has no
internal battery and must remain plugged in to operate.

Dimensions
120mm x 120mm x 35mm.

App requirements
The Solstice Home app requires iOS 15 or later, or Android 10 or later.

Box contents
The box includes: one SH-100 hub, one USB-C cable, one 9W power adapter, a
quick start guide, and a mounting bracket.
"""

PRODUCT_GUIDE_PAGE_2 = """Solstice Home Hub SH-100 -- Product Guide (continued)

Setup steps
1. Plug the hub into power using the included USB-C cable and adapter.
2. Open the Solstice Home app and tap "Add Hub".
3. Scan the QR code printed on the base of the hub.
4. Connect the hub to your Wi-Fi network when prompted.
5. Give the hub a name to finish setup.

LED status reference
Solid blue means the hub is powered on and ready to pair.
Blinking blue means the hub is in pairing mode.
Solid green means the hub is connected and running the latest firmware.
Blinking red means a connectivity error.
Solid red means a firmware update failed.

Firmware updates
The hub checks for firmware updates automatically and installs them
overnight between 2:00 AM and 4:00 AM local time, unless automatic updates
are disabled in the app settings.
"""

# ---------------------------------------------------------------------------
# Support policy (DOCX) -- warranty, returns, Priority Care, support channels
# ---------------------------------------------------------------------------

SUPPORT_POLICY_PARAGRAPHS = [
    ("heading", "Solstice Home Hub SH-100 -- Support Policy"),
    ("heading", "Warranty"),
    (
        "body",
        "The SH-100 is covered by a 12-month limited hardware warranty from the "
        "original date of purchase. The warranty covers manufacturing defects. "
        "It does not cover water damage, physical damage, or units that have "
        "been opened or modified by anyone other than Solstice Devices Inc.",
    ),
    ("heading", "Returns"),
    (
        "body",
        "Unused, undamaged units in their original packaging may be returned "
        "within 30 days of delivery for a full refund. The buyer pays return "
        "shipping unless the unit arrived defective, in which case Solstice "
        "Devices Inc. covers return shipping.",
    ),
    ("heading", "Replacing a defective unit within warranty"),
    (
        "body",
        "To request a replacement for a hardware defect within the 12-month "
        "warranty period, open the Solstice Home app, go to Support, and select "
        "Report an Issue. Ship the defective unit back using the prepaid label "
        "provided in the app. A replacement unit ships within 5 business days "
        "of Solstice Devices Inc. receiving the returned hub.",
    ),
    ("heading", "Out-of-warranty repair"),
    (
        "body",
        "If the 12-month warranty has expired, hardware repair is available for "
        "a flat fee of $39.99. The hub must be shipped in for repair, and "
        "turnaround time is 10 to 14 business days.",
    ),
    ("heading", "Priority Care"),
    (
        "body",
        "Priority Care is an optional paid add-on for $4.99 per month. It "
        "includes phone support, a 2-year extended warranty in place of the "
        "standard 12-month warranty, and priority access to new firmware "
        "before general release.",
    ),
    ("heading", "Support channels"),
    (
        "body",
        "The table below lists how to reach support and what to expect from "
        "each channel.",
    ),
]

SUPPORT_CHANNELS_TABLE = [
    ["Channel", "Availability", "Typical response time"],
    ["Email (support@solsticedevices.example)", "24/7", "24-48 hours"],
    ["Live chat (solsticedevices.example)", "Weekdays, 9 AM-6 PM PT", "Same session"],
    ["Phone", "Priority Care subscribers only", "Immediate, weekdays"],
]

# ---------------------------------------------------------------------------
# Troubleshooting guide (TXT)
# ---------------------------------------------------------------------------

TROUBLESHOOTING_GUIDE = """Solstice Home Hub SH-100 -- Troubleshooting Guide

Hub will not power on
Check that the USB-C cable is fully seated in the hub and the adapter.
Try a different power outlet. If the hub still will not power on, hold the
reset button on the base for 10 seconds and try again.

Blinking red LED (connectivity error)
Confirm your Wi-Fi router has its 2.4GHz or 5GHz band enabled. Move the hub
closer to the router. Restart the router, then re-run setup in the app.

Solid red LED (firmware update failed)
Do not unplug the hub. It will automatically retry the update up to 3 times
over the next 10 minutes. If the LED is still solid red after 3 retries,
factory reset the hub and re-pair it.

Factory reset steps
Hold the reset button on the base of the hub for 15 seconds, until the LED
flashes white. The hub returns to pairing mode. Re-add it in the Solstice
Home app as if it were new.

Persistent red LED after a factory reset
If a blinking red or solid red LED continues after a factory reset, this
indicates a hardware fault rather than a setup problem. If the hub is still
within its 12-month warranty, contact support through the app to request a
replacement at no cost. If the warranty has expired, the standard $39.99
flat-fee repair applies instead.

Devices dropping offline randomly
Check the app for an available firmware update. Confirm you have not
exceeded the maximum of 150 connected devices -- going over this limit can
cause devices to drop offline. If you are near the limit, remove unused
devices or reboot the hub weekly to reduce mesh network congestion.
"""


def write_markdown_sources() -> None:
    SOURCE_DIR.mkdir(exist_ok=True)

    (SOURCE_DIR / "product_guide.md").write_text(
        "# Solstice Home Hub SH-100 -- Product Guide\n\n"
        + PRODUCT_GUIDE_PAGE_1.split("\n", 1)[1].strip()
        + "\n\n---\n\n"
        + PRODUCT_GUIDE_PAGE_2.split("\n", 1)[1].strip()
        + "\n",
        encoding="utf-8",
    )

    lines = ["# Solstice Home Hub SH-100 -- Support Policy\n"]
    for kind, text in SUPPORT_POLICY_PARAGRAPHS[1:]:
        lines.append(f"## {text}\n" if kind == "heading" else f"{text}\n")
    lines.append("\n| " + " | ".join(SUPPORT_CHANNELS_TABLE[0]) + " |")
    lines.append("|" + "|".join(["---"] * len(SUPPORT_CHANNELS_TABLE[0])) + "|")
    for row in SUPPORT_CHANNELS_TABLE[1:]:
        lines.append("| " + " | ".join(row) + " |")
    (SOURCE_DIR / "support_policy.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    (SOURCE_DIR / "troubleshooting_guide.md").write_text(
        "# " + TROUBLESHOOTING_GUIDE.split("\n", 1)[0].replace(" -- ", " — ") + "\n\n"
        + TROUBLESHOOTING_GUIDE.split("\n", 1)[1].strip()
        + "\n",
        encoding="utf-8",
    )


def write_product_guide_pdf() -> None:
    pdf = FPDF()
    for page_text in (PRODUCT_GUIDE_PAGE_1, PRODUCT_GUIDE_PAGE_2):
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(0, 6, page_text)
    pdf.output(str(SAMPLES_DIR / "product_guide.pdf"))


def write_support_policy_docx() -> None:
    document = docx.Document()
    for kind, text in SUPPORT_POLICY_PARAGRAPHS:
        if kind == "heading":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)

    table = document.add_table(rows=len(SUPPORT_CHANNELS_TABLE), cols=3)
    for row_idx, row in enumerate(SUPPORT_CHANNELS_TABLE):
        for col_idx, value in enumerate(row):
            table.rows[row_idx].cells[col_idx].text = value

    document.save(str(SAMPLES_DIR / "support_policy.docx"))


def write_troubleshooting_guide_txt() -> None:
    (SAMPLES_DIR / "troubleshooting_guide.txt").write_text(
        TROUBLESHOOTING_GUIDE, encoding="utf-8"
    )


if __name__ == "__main__":
    write_markdown_sources()
    write_product_guide_pdf()
    write_support_policy_docx()
    write_troubleshooting_guide_txt()
    print("Generated sample documents in", SAMPLES_DIR)
