from pathlib import Path

import docx
import pytest
from fpdf import FPDF

from app.ingestion.extraction import ExtractionError, extract_units


def test_extract_txt_returns_line_units(tmp_path: Path):
    path = tmp_path / "sample.txt"
    path.write_text("first line\n\nsecond line\n", encoding="utf-8")

    units = extract_units(path, ".txt")

    assert [u.text for u in units] == ["first line", "second line"]
    assert units[0].source == {"kind": "line", "line": 1}
    assert units[1].source == {"kind": "line", "line": 3}


def test_extract_txt_rejects_non_utf8(tmp_path: Path):
    path = tmp_path / "bad.txt"
    path.write_bytes(b"\xff\xfe\x00bad")

    with pytest.raises(ExtractionError):
        extract_units(path, ".txt")


def test_extract_txt_rejects_empty_file(tmp_path: Path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n", encoding="utf-8")

    with pytest.raises(ExtractionError):
        extract_units(path, ".txt")


def test_extract_pdf_returns_page_units(tmp_path: Path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Hello from page one")
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Hello from page two")
    path = tmp_path / "sample.pdf"
    pdf.output(str(path))

    units = extract_units(path, ".pdf")

    assert len(units) == 2
    assert "page one" in units[0].text
    assert units[0].source == {"kind": "page", "page": 1}
    assert "page two" in units[1].text
    assert units[1].source == {"kind": "page", "page": 2}


def test_extract_pdf_rejects_scanned_empty_pdf(tmp_path: Path):
    pdf = FPDF()
    pdf.add_page()
    path = tmp_path / "blank.pdf"
    pdf.output(str(path))

    with pytest.raises(ExtractionError):
        extract_units(path, ".pdf")


def test_extract_docx_returns_paragraph_and_table_units(tmp_path: Path):
    document = docx.Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell a"
    table.rows[0].cells[1].text = "cell b"
    path = tmp_path / "sample.docx"
    document.save(str(path))

    units = extract_units(path, ".docx")

    paragraph_units = [u for u in units if u.source["kind"] == "paragraph"]
    table_units = [u for u in units if u.source["kind"] == "table_row"]

    assert [u.text for u in paragraph_units] == ["First paragraph", "Second paragraph"]
    assert table_units[0].text == "cell a | cell b"
    assert table_units[0].source == {"kind": "table_row", "table": 0, "row": 0}


def test_extract_docx_rejects_empty_document(tmp_path: Path):
    document = docx.Document()
    path = tmp_path / "empty.docx"
    document.save(str(path))

    with pytest.raises(ExtractionError):
        extract_units(path, ".docx")


def test_extract_unsupported_extension(tmp_path: Path):
    path = tmp_path / "sample.exe"
    path.write_bytes(b"binary")

    with pytest.raises(ExtractionError):
        extract_units(path, ".exe")
